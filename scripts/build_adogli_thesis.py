"""
Générateur Complet Haute Fidélité du Mémoire Master SUNU Bank Togo
Conforme à 100% à la présentation du Mémoire ADOGLI (Collège de Paris / Université de Lomé)
et intégrant les captures d'écran ORIGINALES de l'application.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_OUTPUT = "MEMOIRE_MASTER_SUNU_BANK_TOGO_RAG_FINAL.docx"

# Palette de couleurs officielles
C_NAVY = RGBColor(0x00, 0x33, 0x66)       # Bleu Nuit Titres / En-têtes (#003366)
C_RED = RGBColor(0xE2, 0x1E, 0x26)        # Rouge Institutionnel SUNU (#E21E26)
C_DARK = RGBColor(0x1A, 0x20, 0x2C)       # Texte principal (#1A202C)
C_MUTED = RGBColor(0x4A, 0x55, 0x68)      # Gris sous-titres / sources (#4A5568)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Définit le padding interne d'une cellule."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>\n'
                      f'<w:top w:w="{top}" w:type="dxa"/>\n'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>\n'
                      f'<w:left w:w="{left}" w:type="dxa"/>\n'
                      f'<w:right w:w="{right}" w:type="dxa"/>\n'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """Définit les bordures d'une cellule XML."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>\n'
                          f'<w:top w:val="{kwargs.get("top", "none")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="{kwargs.get("color", "CBD5E1")}"/>\n'
                          f'<w:left w:val="{kwargs.get("left", "none")}" w:sz="{kwargs.get("left_sz", "4")}" w:space="0" w:color="{kwargs.get("color", "CBD5E1")}"/>\n'
                          f'<w:bottom w:val="{kwargs.get("bottom", "none")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="{kwargs.get("color", "CBD5E1")}"/>\n'
                          f'<w:right w:val="{kwargs.get("right", "none")}" w:sz="{kwargs.get("right_sz", "4")}" w:space="0" w:color="{kwargs.get("color", "CBD5E1")}"/>\n'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

def add_header_footer_lines(section, is_roman=False):
    """Configure l'en-tête et le pied de page exactement comme Adogli."""
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = hp.add_run("CONCEPTION D'UN ASSISTANT CONVERSATIONNEL INTELLIGENT BASÉ SUR LE RAG POUR L'ACCOMPAGNEMENT PRÉCONTRACTUEL EN BANCASSURANCE : CAS DE SUNU BANK TOGO")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Bordure inférieure sous l'en-tête
    pPr = hp._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}>\n'
                     f'<w:bottom w:val="single" w:sz="6" w:space="4" w:color="A0AEC0"/>\n'
                     f'</w:pBdr>')
    pPr.append(pBdr)
    
    # Pied de page
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    
    # Bordure supérieure au-dessus du pied de page
    fpPr = fp._p.get_or_add_pPr()
    fpBdr = parse_xml(f'<w:pBdr {nsdecls("w")}>\n'
                      f'<w:top w:val="single" w:sz="6" w:space="4" w:color="A0AEC0"/>\n'
                      f'</w:pBdr>')
    fpPr.append(fpBdr)
    
    # Run Auteur à gauche
    run_left = fp.add_run("Rédigé et présenté par ADOGLI Jean-Paul")
    run_left.font.name = "Calibri"
    run_left.font.size = Pt(9)
    run_left.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    
    # Tabulation vers la droite
    run_tab = fp.add_run("\t\t")
    
    # Champ de numéro de page
    run_page = fp.add_run()
    run_page.font.name = "Calibri"
    run_page.font.size = Pt(9.5)
    run_page.font.bold = True
    run_page.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run_page._r.append(fldSimple)

def add_callout(doc, text, title=""):
    """Crée un encadré académique avec bordure gauche épaisse et fond teinté."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}>\n'
                     f'<w:left w:val="single" w:sz="24" w:space="12" w:color="003366"/>\n'
                     f'<w:top w:val="none"/>\n'
                     f'<w:right w:val="none"/>\n'
                     f'<w:bottom w:val="none"/>\n'
                     f'</w:pBdr>')
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    pPr.append(pBdr)
    pPr.append(shd)
    
    if title:
        rt = p.add_run(f"• {title}\n")
        rt.font.name = "Calibri"
        rt.font.size = Pt(10)
        rt.font.bold = True
        rt.font.color.rgb = C_NAVY
        
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = C_DARK
    return p

def add_figure_with_source(doc, img_path, figure_title, source_text="Source : Auteur (2026)", width_in=6.0):
    """Insère une figure avec sa légende et sa source conformes à Adogli."""
    if not os.path.exists(img_path):
        print(f"[ATTENTION] Image introuvable : {img_path}")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(6)
    p_img.add_run().add_picture(img_path, width=Inches(width_in))
    
    # Titre de la figure (Italique, Adogli style)
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(2)
    r_cap = p_cap.add_run(figure_title)
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = C_NAVY
    
    # Source
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_src.paragraph_format.space_before = Pt(0)
    p_src.paragraph_format.space_after = Pt(14)
    r_src = p_src.add_run(source_text)
    r_src.font.name = "Calibri"
    r_src.font.size = Pt(9)
    r_src.font.italic = True
    r_src.font.color.rgb = C_MUTED

def parse_markdown_tables(doc, table_lines, source_text="Source : Auteur (2026)"):
    """Parse et insère un tableau Markdown stylisé en tableau Word professionnel."""
    if not table_lines:
        return
    rows_data = []
    for line in table_lines:
        if re.match(r"^\|(\s*[-:]+\s*\|)+$", line.strip()):
            continue # Ligne séparatrice Markdown
        cols = [c.strip() for c in line.strip().split('|')[1:-1]]
        if cols:
            rows_data.append(cols)
            
    if not rows_data:
        return
        
    num_cols = len(rows_data[0])
    num_rows = len(rows_data)
    
    tbl = doc.add_table(rows=num_rows, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_w = Inches(6.2 / num_cols)
    for row_idx, row in enumerate(rows_data):
        is_header = (row_idx == 0)
        bg_color = "003366" if is_header else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        for col_idx, text in enumerate(row):
            if col_idx < num_cols:
                cell = tbl.cell(row_idx, col_idx)
                cell.width = col_w
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                tcPr = cell._tc.get_or_add_tcPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
                tcPr.append(shd)
                
                if is_header:
                    set_cell_border(cell, top="single", bottom="double", color="003366", bottom_sz="12")
                else:
                    set_cell_border(cell, bottom="single", color="E2E8F0", bottom_sz="4")
                    
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                clean_text = text.replace("**", "").replace("*", "")
                r = p.add_run(clean_text)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5 if is_header else 9)
                if is_header or "**" in text:
                    r.font.bold = True
                r.font.color.rgb = C_WHITE if is_header else C_DARK

    # Source du tableau
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_src.paragraph_format.space_before = Pt(2)
    p_src.paragraph_format.space_after = Pt(12)
    r_src = p_src.add_run(source_text)
    r_src.font.name = "Calibri"
    r_src.font.size = Pt(8.5)
    r_src.font.italic = True
    r_src.font.color.rgb = C_MUTED

def parse_markdown_section_content(doc, md_text):
    """Parse un chapitre Markdown et insère titres, paragraphes, puces, encadrés et tableaux."""
    lines = md_text.split('\n')
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        line_s = line.strip()
        
        if not line_s or line_s.startswith("# "):
            idx += 1
            continue
            
        # Détection d'un tableau Markdown
        if line_s.startswith('|'):
            tbl_lines = []
            while idx < len(lines) and lines[idx].strip().startswith('|'):
                tbl_lines.append(lines[idx].strip())
                idx += 1
            parse_markdown_tables(doc, tbl_lines)
            continue
            
        # Titre Niveau 2
        if line_s.startswith("## ") or re.match(r"^[I|V|X\d]+\.\d+\.\s+", line_s):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            clean_title = re.sub(r"^##\s+", "", line_s)
            r = p.add_run(clean_title)
            r.font.name = "Calibri"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = C_NAVY
            idx += 1
            continue
            
        # Titre Niveau 3
        if line_s.startswith("### ") or re.match(r"^[I|V|X\d]+\.\d+\.\d+\.\s+", line_s):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            clean_title = re.sub(r"^###\s+", "", line_s)
            r = p.add_run(clean_title)
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = C_DARK
            idx += 1
            continue
            
        # Puces
        if line_s.startswith("- ") or line_s.startswith("* ") or line_s.startswith("• "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            clean_txt = re.sub(r"^[-*•]\s+", "", line_s)
            r = p.add_run("• " + clean_txt)
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            idx += 1
            continue
            
        # Encadrés réglementaires CIMA / IPDCP
        if "[Encadré réglementaire]" in line_s or "Article 74" in line_s or "Article 6 du Code CIMA" in line_s:
            add_callout(doc, line_s, "Cadre Réglementaire Obligatoire (CIMA & IPDCP Togo)")
            idx += 1
            continue
            
        # Paragraphe standard
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(line_s)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        idx += 1

def build_complete_thesis():
    """Génère le document Word complet intégrant les 12 figures, tables et styles Adogli."""
    doc = docx.Document()
    
    # Marges normalisées
    for s in doc.sections:
        s.top_margin = Inches(0.98)
        s.bottom_margin = Inches(0.98)
        s.left_margin = Inches(0.98)
        s.right_margin = Inches(0.98)
        s.different_first_page_header_footer = True
        
    print("Étape 1 : Construction de la page de couverture...")
    from cover_adogli import create_cover_page
    create_cover_page(doc)
    doc.add_page_break()
    
    print("Étape 2 : Configuration des pages préliminaires...")
    sec_prelim = doc.sections[0]
    add_header_footer_lines(sec_prelim, is_roman=True)
    
    # DÉDICACE
    p_ded_t = doc.add_paragraph()
    p_ded_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ded_t.paragraph_format.space_before = Pt(18)
    p_ded_t.paragraph_format.space_after = Pt(120)
    r_ded_t = p_ded_t.add_run("DEDICACE")
    r_ded_t.font.name = "Calibri"
    r_ded_t.font.size = Pt(13)
    r_ded_t.font.bold = True
    
    p_ded_c = doc.add_paragraph()
    p_ded_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ded_c.paragraph_format.space_after = Pt(14)
    r_ded_c = p_ded_c.add_run("À MES PARENTS ET À MA FAMILLE")
    r_ded_c.font.name = "Calibri"
    r_ded_c.font.size = Pt(12)
    r_ded_c.font.bold = True
    
    p_ded_sub = doc.add_paragraph()
    p_ded_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ded_sub = p_ded_sub.add_run("Pour leur soutien indéfectible, leurs sacrifices et leurs encouragements constants.")
    r_ded_sub.font.name = "Calibri"
    r_ded_sub.font.size = Pt(11)
    r_ded_sub.font.italic = True
    doc.add_page_break()
    
    # REMERCIEMENTS
    p_rem_t = doc.add_paragraph()
    p_rem_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rem_t.paragraph_format.space_before = Pt(18)
    p_rem_t.paragraph_format.space_after = Pt(24)
    r_rem_t = p_rem_t.add_run("REMERCIEMENTS")
    r_rem_t.font.name = "Calibri"
    r_rem_t.font.size = Pt(13)
    r_rem_t.font.bold = True
    
    remerciements_text = [
        "Aucun travail de recherche ne saurait être l'œuvre d'une seule personne. Sa réalisation résulte toujours du concours de plusieurs acteurs, qu'ils soient académiques, professionnels, familiaux ou amicaux. C'est pourquoi nous tenons à exprimer notre profonde gratitude à toutes celles et ceux qui, de près ou de loin, ont contribué à l'aboutissement de ce mémoire.",
        "Nos remerciements s'adressent particulièrement à :",
        "• Monsieur Latevi Sena LAWSON, notre encadreur académique, pour sa disponibilité, sa rigueur scientifique et ses précieux conseils tout au long de cette recherche ;",
        "• Monsieur Kokou AGBOKOU, notre encadreur professionnel, pour son accompagnement technique et sa bienveillance au sein de la structure d'accueil ;",
        "• La Direction et l'ensemble des équipes de SUNU Bank Togo, particulièrement le pôle Bancassurance, pour leur accueil chaleureux et les données précieuses mises à notre disposition ;",
        "• Le corps professoral et l'administration du Collège de Paris Supérieur et de l'Université de Lomé, pour la qualité de l'enseignement dispensé et l'excellence de notre formation ;",
        "• Nos chers parents, frères, sœurs et amis, pour leurs prières, leurs encouragements constants et leur patience inaltérable tout au long de ce parcours académique."
    ]
    for para in remerciements_text:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(para)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    doc.add_page_break()
    
    # RÉSUMÉ
    p_res_t = doc.add_paragraph()
    p_res_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_res_t.paragraph_format.space_before = Pt(18)
    p_res_t.paragraph_format.space_after = Pt(24)
    r_res_t = p_res_t.add_run("RESUME")
    r_res_t.font.name = "Calibri"
    r_res_t.font.size = Pt(13)
    r_res_t.font.bold = True
    
    resume_text = [
        "Dans le secteur financier ouest-africain, la bancassurance constitue un canal stratégique de distribution des produits d'assurance vie. Cependant, la complexité des clauses contractuelles et le niveau limité de littératie financière créent d'importantes asymétries d'information chez les prospects, amplifiées par la saturation des conseillers bancaires en agence.",
        "Le présent mémoire étudie la conception, l'implémentation et l'évaluation d'un assistant conversationnel intelligent fondé sur l'architecture RAG (Retrieval-Augmented Generation), appliqué au portefeuille d'assurance vie de SUNU Bank Togo (offres Visa Études, Visa Études Plus et Horizon Retraite).",
        "S'appuyant sur une démarche mixte et un corpus documentaire rigoureusement calibré (150 chunks indexés sous ChromaDB avec embeddings all-MiniLM-L6-v2), la solution intègre des mécanismes stricts de vérification de fidélité documentaire (RAGAS), de citation des clauses contractuelles et de refus hors périmètre avec escalade humaine, en conformité avec le Code des Assurances CIMA (Articles 6 et 74) et la Loi togolaise n° 2019-014 sur la protection des données personnelles.",
        "Les évaluations empiriques sur un jeu de test de 75 requêtes démontrent un taux de réussite de recherche (Hit@5) de 78,7 %, un MRR de 0,434 et une fidélité factuelle RAGAS de 0,84, confirmant la viabilité opérationnelle de l'intelligence artificielle générative pour démocratiser l'information précontractuelle en milieu bancaire.",
        "Mots-clés : Bancassurance, RAG, Intelligence Artificielle Générative, Assurance Vie, Information Précontractuelle, Code CIMA, SUNU Bank Togo."
    ]
    for para in resume_text:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(para)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        if para.startswith("Mots-clés"):
            r.font.bold = True
    doc.add_page_break()
    
    # ABSTRACT
    p_abs_t = doc.add_paragraph()
    p_abs_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_t.paragraph_format.space_before = Pt(18)
    p_abs_t.paragraph_format.space_after = Pt(24)
    r_abs_t = p_abs_t.add_run("ABSTRACT")
    r_abs_t.font.name = "Calibri"
    r_abs_t.font.size = Pt(13)
    r_abs_t.font.bold = True
    
    abstract_text = [
        "In the West African financial sector, bancassurance represents a vital strategic channel for distributing life insurance policies. However, the intricacies of contract clauses and modest financial literacy levels among prospects often lead to informational asymmetries, exacerbated by branch advisors' heavy workload.",
        "This master's thesis investigates the design, implementation, and empirical validation of an intelligent conversational assistant based on Retrieval-Augmented Generation (RAG) architecture, specifically developed for the life insurance portfolio of SUNU Bank Togo (Visa Études, Visa Études Plus, and Horizon Retraite).",
        "Adopting a mixed-method research framework and a calibrated corporate corpus (150 chunks indexed in ChromaDB using all-MiniLM-L6-v2 embeddings), the platform embeds strict documentary grounding, dynamic clause citations, out-of-scope refusal triggers, and human advisor escalation, ensuring full regulatory compliance with CIMA Insurance Code (Articles 6 & 74) and Togolese Data Protection Law No. 2019-014.",
        "Empirical benchmarks across 75 test queries reveal a 78.7% Hit@5 retrieval rate, an MRR of 0.434, and an average RAGAS faithfulness score of 0.84, validating the operational feasibility of generative AI in democratizing precontractual financial information.",
        "Keywords : Bancassurance, RAG, Generative AI, Life Insurance, Precontractual Information, CIMA Code, SUNU Bank Togo."
    ]
    for para in abstract_text:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(para)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        if para.startswith("Keywords"):
            r.font.bold = True
    doc.add_page_break()
    
    # SOMMAIRE
    p_som_t = doc.add_paragraph()
    p_som_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_som_t.paragraph_format.space_before = Pt(18)
    p_som_t.paragraph_format.space_after = Pt(24)
    r_som_t = p_som_t.add_run("SOMMAIRE")
    r_som_t.font.name = "Calibri"
    r_som_t.font.size = Pt(13)
    r_som_t.font.bold = True
    
    sommaire_items = [
        ("DEDICACE", "I"),
        ("REMERCIEMENTS", "II"),
        ("RESUME", "III"),
        ("ABSTRACT", "IV"),
        ("SOMMAIRE", "V"),
        ("LISTE DES TABLEAUX", "VI"),
        ("LISTE DES FIGURES ET GRAPHIQUES", "VII"),
        ("LISTE DES ABREVIATIONS", "VIII"),
        ("INTRODUCTION GENERALE", "1"),
        ("1. CONTEXTE GENERAL DE L'ETUDE", "2"),
        ("2. PROBLEMATIQUE DE L'ETUDE", "3"),
        ("3. HYPOTHESES DE L'ETUDE", "4"),
        ("4. OBJECTIFS DE L'ETUDE", "5"),
        ("5. JUSTIFICATION DE L'ETUDE", "6"),
        ("6. DELIMITATION DE L'ETUDE", "7"),
        ("7. PLAN DU MEMOIRE", "8"),
        ("CHAPITRE I : CADRE THEORIQUE ET CONCEPTUEL", "9"),
        ("CHAPITRE II : METHODOLOGIE DE L'ETUDE", "22"),
        ("CHAPITRE III : PRESENTATION DU CADRE D'ETUDE ET DU CORPUS", "33"),
        ("CHAPITRE IV : CONCEPTION, IMPLEMENTATION ET EVALUATION DU PROTOTYPE RAG", "42"),
        ("CONCLUSION GENERALE", "68"),
        ("BIBLIOGRAPHIE ET WEBGRAPHIE", "72"),
        ("ANNEXES", "76"),
        ("TABLE DES MATIERES", "79")
    ]
    for label, page_no in sommaire_items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        dots_count = max(4, 75 - len(label) - len(page_no))
        dots_str = " " + "." * dots_count + " "
        r_lbl = p.add_run(label)
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(10.5)
        if "CHAPITRE" in label or "INTRODUCTION" in label or "CONCLUSION" in label or "SOMMAIRE" in label:
            r_lbl.font.bold = True
        r_dots = p.add_run(dots_str)
        r_dots.font.name = "Calibri"
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = C_MUTED
        r_pg = p.add_run(page_no)
        r_pg.font.name = "Calibri"
        r_pg.font.size = Pt(10.5)
        r_pg.font.bold = True
    doc.add_page_break()
    
    # LISTE DES TABLEAUX
    p_lt_t = doc.add_paragraph()
    p_lt_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lt_t.paragraph_format.space_before = Pt(18)
    p_lt_t.paragraph_format.space_after = Pt(24)
    r_lt_t = p_lt_t.add_run("LISTE DES TABLEAUX")
    r_lt_t.font.name = "Calibri"
    r_lt_t.font.size = Pt(13)
    r_lt_t.font.bold = True
    
    tableaux_list = [
        ("Tableau I.1 : Comparatif des canaux d'information bancassurance", "13"),
        ("Tableau I.2 : Synthèse des cadres réglementaires (CIMA, UEMOA, IPDCP)", "18"),
        ("Tableau II.1 : Dimensions conceptuelles et indicateurs d'évaluation", "27"),
        ("Tableau IV.1 : Résultats de la recherche documentaire (75 questions)", "48"),
        ("Tableau IV.2 : Expérience de calibration du chunking (6 configurations)", "51"),
        ("Tableau IV.3 : Expérience de comparaison des modèles d'embeddings", "53"),
        ("Tableau IV.4 : Performances de retrieval détaillées par catégorie de requêtes", "55"),
        ("Tableau IV.5 : Résultats de l'évaluation RAGAS sur le prototype", "58"),
        ("Tableau IV.6 : Grille d'évaluation experte de conformité métier", "61"),
        ("Tableau IV.7 : Synthèse de la validation des hypothèses de recherche", "64")
    ]
    for label, page_no in tableaux_list:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        dots_count = max(4, 75 - len(label) - len(page_no))
        dots_str = " " + "." * dots_count + " "
        r_lbl = p.add_run(label)
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(10)
        r_dots = p.add_run(dots_str)
        r_dots.font.name = "Calibri"
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = C_MUTED
        r_pg = p.add_run(page_no)
        r_pg.font.name = "Calibri"
        r_pg.font.size = Pt(10)
        r_pg.font.bold = True
    doc.add_page_break()
    
    # LISTE DES FIGURES ET GRAPHIQUES
    p_lf_t = doc.add_paragraph()
    p_lf_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lf_t.paragraph_format.space_before = Pt(18)
    p_lf_t.paragraph_format.space_after = Pt(24)
    r_lf_t = p_lf_t.add_run("LISTE DES FIGURES ET GRAPHIQUES")
    r_lf_t.font.name = "Calibri"
    r_lf_t.font.size = Pt(13)
    r_lf_t.font.bold = True
    
    figures_list = [
        ("Figure I.1 : Architecture générale d'un système RAG appliqué à l'assurance vie", "15"),
        ("Figure I.2 : Pipeline séquentiel de traitement et de contextualisation d'une requête", "17"),
        ("Figure III.1 : Structure et volumétrie du corpus documentaire SUNU Bank (150 chunks)", "38"),
        ("Figure IV.1 : Répartition des 75 questions du jeu de test par catégorie", "46"),
        ("Figure IV.2 : Évaluation empirique de la calibration du chunking", "50"),
        ("Figure IV.3 : Comparaison des performances des modèles d'embeddings", "52"),
        ("Figure IV.4 : Performances de recherche par catégorie de questions", "54"),
        ("Figure IV.5 : Interface d'accueil du Concierge Financier SUNU Bank (Mode Sombre)", "62"),
        ("Figure IV.6 : Interface d'accueil du Concierge Financier (Mode Clair)", "63"),
        ("Figure IV.7 : Interface conversationnelle interactive de la Discussion IA", "64"),
        ("Figure IV.8 : Réponse en temps réel avec citation des sources (Visa Études)", "65"),
        ("Figure IV.9 : Écran officiel d'authentification et de gestion de session", "66")
    ]
    for label, page_no in figures_list:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        dots_count = max(4, 75 - len(label) - len(page_no))
        dots_str = " " + "." * dots_count + " "
        r_lbl = p.add_run(label)
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(10)
        r_dots = p.add_run(dots_str)
        r_dots.font.name = "Calibri"
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = C_MUTED
        r_pg = p.add_run(page_no)
        r_pg.font.name = "Calibri"
        r_pg.font.size = Pt(10)
        r_pg.font.bold = True
    doc.add_page_break()
    
    # LISTE DES ABRÉVIATIONS
    p_la_t = doc.add_paragraph()
    p_la_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_la_t.paragraph_format.space_before = Pt(18)
    p_la_t.paragraph_format.space_after = Pt(24)
    r_la_t = p_la_t.add_run("LISTE DES ABREVIATIONS")
    r_la_t.font.name = "Calibri"
    r_la_t.font.size = Pt(13)
    r_la_t.font.bold = True
    
    abrevs = [
        ("API", "Application Programming Interface (Interface de Programmation d'Applications)"),
        ("BCEAO", "Banque Centrale des États de l'Afrique de l'Ouest"),
        ("BPEC", "Banque Populaire pour l'Épargne et le Crédit (devenue SUNU Bank Togo)"),
        ("CIMA", "Conférence Interafricaine des Marchés d'Assurances"),
        ("CRAC", "Commission Régionale de Contrôle des Assurances"),
        ("FANAF", "Fédération des Sociétés d'Assurances de Droit National Africaines"),
        ("FAQ", "Foire Aux Questions"),
        ("IA", "Intelligence Artificielle"),
        ("IPDCP", "Instance de Protection des Données à Caractère Personnel (Togo)"),
        ("LLM", "Large Language Model (Modèle de Langage à Grande Échelle)"),
        ("MRR", "Mean Reciprocal Rank (Rang Moyen Réciproque)"),
        ("NLP", "Natural Language Processing (Traitement Automatique du Langage Naturel)"),
        ("RAG", "Retrieval-Augmented Generation (Génération Augmentée par Récupération)"),
        ("RAGAS", "Retrieval Augmented Generation Assessment"),
        ("TAM", "Technology Acceptance Model (Modèle d'Acceptation de la Technologie)"),
        ("UEMOA", "Union Économique et Monétaire Ouest-Africaine")
    ]
    for abbr, exp in abrevs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r_ab = p.add_run(f"{abbr} : ")
        r_ab.font.name = "Calibri"
        r_ab.font.size = Pt(10)
        r_ab.font.bold = True
        r_ex = p.add_run(exp)
        r_ex.font.name = "Calibri"
        r_ex.font.size = Pt(10)
    doc.add_page_break()
    
    print("Étape 3 : Construction du Corps Principal...")
    sec_main = doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_header_footer_lines(sec_main, is_roman=False)
    
    # ----------------------------------------------------
    # INTRODUCTION GÉNÉRALE
    # ----------------------------------------------------
    p_ig_t = doc.add_paragraph()
    p_ig_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ig_t.paragraph_format.space_before = Pt(24)
    p_ig_t.paragraph_format.space_after = Pt(24)
    r_ig_t = p_ig_t.add_run("INTRODUCTION GENERALE")
    r_ig_t.font.name = "Calibri"
    r_ig_t.font.size = Pt(15)
    r_ig_t.font.bold = True
    r_ig_t.font.color.rgb = C_NAVY
    
    with open("Memoire/01_introduction_generale.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
    doc.add_page_break()
    
    # ----------------------------------------------------
    # CHAPITRE I : CADRE THÉORIQUE ET CONCEPTUEL
    # ----------------------------------------------------
    p_c1_t = doc.add_paragraph()
    p_c1_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c1_t.paragraph_format.space_before = Pt(24)
    p_c1_t.paragraph_format.space_after = Pt(24)
    r_c1_t = p_c1_t.add_run("CHAPITRE I : CADRE THEORIQUE ET CONCEPTUEL")
    r_c1_t.font.name = "Calibri"
    r_c1_t.font.size = Pt(14)
    r_c1_t.font.bold = True
    r_c1_t.font.color.rgb = C_NAVY
    
    with open("Memoire/02_chapitre1_cadre_theorique.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
        
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_architecture.png",
        "Figure I.1 : Architecture générale d'un système RAG appliqué à l'assurance vie",
        "Source : Conception auteur adaptée de Lewis et al. (2020)"
    )
    
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_pipeline.png",
        "Figure I.2 : Pipeline séquentiel de traitement et de contextualisation d'une requête",
        "Source : Conception auteur (2026)"
    )
    doc.add_page_break()
    
    # ----------------------------------------------------
    # CHAPITRE II : MÉTHODOLOGIE DE L'ÉTUDE
    # ----------------------------------------------------
    p_c2_t = doc.add_paragraph()
    p_c2_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c2_t.paragraph_format.space_before = Pt(24)
    p_c2_t.paragraph_format.space_after = Pt(24)
    r_c2_t = p_c2_t.add_run("CHAPITRE II : METHODOLOGIE DE L'ETUDE")
    r_c2_t.font.name = "Calibri"
    r_c2_t.font.size = Pt(14)
    r_c2_t.font.bold = True
    r_c2_t.font.color.rgb = C_NAVY
    
    with open("Memoire/03_chapitre2_methodologie.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
    doc.add_page_break()
    
    # ----------------------------------------------------
    # CHAPITRE III : PRÉSENTATION DU TERRAIN D'ÉTUDE ET DU CORPUS
    # ----------------------------------------------------
    p_c3_t = doc.add_paragraph()
    p_c3_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c3_t.paragraph_format.space_before = Pt(24)
    p_c3_t.paragraph_format.space_after = Pt(24)
    r_c3_t = p_c3_t.add_run("CHAPITRE III : PRESENTATION DU CADRE D'ETUDE ET DU CORPUS")
    r_c3_t.font.name = "Calibri"
    r_c3_t.font.size = Pt(14)
    r_c3_t.font.bold = True
    r_c3_t.font.color.rgb = C_NAVY
    
    with open("Memoire/04_chapitre3_terrain.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
        
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_corpus_distribution.png",
        "Figure III.1 : Structure et volumétrie du corpus documentaire SUNU Bank (150 chunks)",
        "Source : Analyse du corpus documentaire par l'auteur (2026)"
    )
    doc.add_page_break()
    
    # ----------------------------------------------------
    # CHAPITRE IV : CONCEPTION, IMPLÉMENTATION ET ÉVALUATION DU PROTOTYPE RAG
    # ----------------------------------------------------
    p_c4_t = doc.add_paragraph()
    p_c4_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c4_t.paragraph_format.space_before = Pt(24)
    p_c4_t.paragraph_format.space_after = Pt(24)
    r_c4_t = p_c4_t.add_run("CHAPITRE IV : CONCEPTION, IMPLEMENTATION ET EVALUATION DU PROTOTYPE RAG")
    r_c4_t.font.name = "Calibri"
    r_c4_t.font.size = Pt(14)
    r_c4_t.font.bold = True
    r_c4_t.font.color.rgb = C_NAVY
    
    with open("Memoire/05_chapitre4_conception_implementation.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
        
    print("Étape 4 : Insertion des Figures d'Évaluation & VRAIES Captures d'Application...")
    
    # 1. Figures expérimentales
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_categories.png",
        "Figure IV.1 : Répartition des 75 questions du jeu de test selon les thématiques précontractuelles",
        "Source : Conception du jeu de test par l'auteur (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_chunking.png",
        "Figure IV.2 : Évaluation empirique de la calibration du chunking (Taille / Chevauchement)",
        "Source : Expérimentations empiriques réalisées par l'auteur (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_embeddings.png",
        "Figure IV.3 : Comparaison des performances des modèles d'embeddings MiniLM et Multilingue",
        "Source : Expérimentations empiriques réalisées par l'auteur (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Memoire/figures/fig_retrieval_by_category.png",
        "Figure IV.4 : Taux de réussite de la recherche documentaire (Source Hit@5) par catégorie de questions",
        "Source : Analyse des résultats de retrieval par l'auteur (2026)"
    )
    
    # 2. VRAIES CAPTURES ORIGINALES DE L'APPLICATION (Projet/docs/screenshots/)
    add_figure_with_source(
        doc,
        "Projet/docs/screenshots/01_accueil_concierge_sombre.png",
        "Figure IV.5 : Interface d'accueil du Concierge Financier SUNU Bank (Mode Sombre)",
        "Source : Capture originale de l'application prototype SUNU Bank (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Projet/docs/screenshots/02_accueil_concierge_clair.png",
        "Figure IV.6 : Interface d'accueil du Concierge Financier (Mode Clair)",
        "Source : Capture originale de l'application prototype SUNU Bank (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Projet/docs/screenshots/03_chat_discussion_ia.png",
        "Figure IV.7 : Interface conversationnelle interactive de la Discussion IA",
        "Source : Capture originale de l'application prototype SUNU Bank (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Projet/docs/screenshots/04_reponse_concierge_visa_etudes.png",
        "Figure IV.8 : Réponse en temps réel avec citation des sources contractuelles (Visa Études)",
        "Source : Capture originale de l'application prototype SUNU Bank (2026)"
    )
    
    add_figure_with_source(
        doc,
        "Projet/docs/screenshots/05_page_connexion.png",
        "Figure IV.9 : Écran officiel d'authentification et de sécurité des sessions",
        "Source : Capture originale de l'application prototype SUNU Bank (2026)"
    )
    doc.add_page_break()
    
    # ----------------------------------------------------
    # CONCLUSION GÉNÉRALE
    # ----------------------------------------------------
    p_cg_t = doc.add_paragraph()
    p_cg_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cg_t.paragraph_format.space_before = Pt(24)
    p_cg_t.paragraph_format.space_after = Pt(24)
    r_cg_t = p_cg_t.add_run("CONCLUSION GENERALE")
    r_cg_t.font.name = "Calibri"
    r_cg_t.font.size = Pt(14)
    r_cg_t.font.bold = True
    r_cg_t.font.color.rgb = C_NAVY
    
    with open("Memoire/07_conclusion_generale.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
    doc.add_page_break()
    
    # ----------------------------------------------------
    # BIBLIOGRAPHIE ET WEBGRAPHIE
    # ----------------------------------------------------
    p_bib_t = doc.add_paragraph()
    p_bib_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bib_t.paragraph_format.space_before = Pt(24)
    p_bib_t.paragraph_format.space_after = Pt(24)
    r_bib_t = p_bib_t.add_run("BIBLIOGRAPHIE ET WEBGRAPHIE")
    r_bib_t.font.name = "Calibri"
    r_bib_t.font.size = Pt(14)
    r_bib_t.font.bold = True
    r_bib_t.font.color.rgb = C_NAVY
    
    with open("Memoire/08_bibliographie.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())
    doc.add_page_break()
    
    # ----------------------------------------------------
    # ANNEXES
    # ----------------------------------------------------
    p_ann_t = doc.add_paragraph()
    p_ann_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ann_t.paragraph_format.space_before = Pt(24)
    p_ann_t.paragraph_format.space_after = Pt(24)
    r_ann_t = p_ann_t.add_run("ANNEXES")
    r_ann_t.font.name = "Calibri"
    r_ann_t.font.size = Pt(14)
    r_ann_t.font.bold = True
    r_ann_t.font.color.rgb = C_NAVY
    
    with open("Memoire/09_annexes.md", "r", encoding="utf-8") as f:
        parse_markdown_section_content(doc, f.read())

    doc.save(DOC_OUTPUT)
    print(f"=== [TERMINÉ AVEC SUCCÈS] Document généré : {DOC_OUTPUT} ===")

if __name__ == "__main__":
    build_complete_thesis()
