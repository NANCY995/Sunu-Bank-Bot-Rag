"""
Générateur de la Présentation PowerPoint Professionnelle du Mémoire Master SUNU Bank Togo
Inspiré à 100% de la structure, des codes visuels et du design de la présentation de référence (Présentation Emmanuelle.pptx)
et adapté intégralement au sujet RAG & Bancassurance Vie SUNU Bank Togo.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

PPTX_OUTPUT = "PRESENTATION_SOUTENANCE_SUNU_BANK_TOGO_RAG.pptx"

# Palette de couleurs officielles
C_NAVY = RGBColor(0x00, 0x33, 0x66)       # Bleu Nuit Titres / En-têtes (#003366)
C_RED = RGBColor(0xE2, 0x1E, 0x26)        # Rouge Institutionnel SUNU (#E21E26)
C_DARK = RGBColor(0x1A, 0x20, 0x2C)       # Texte principal (#1A202C)
C_MUTED = RGBColor(0x71, 0x80, 0x96)      # Gris sous-titres / sources (#718096)
C_LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)   # Fond de carte (#F8FAFC)
C_BORDER = RGBColor(0xCB, 0xD5, 0xE1)     # Bordure douce (#CBD5E1)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_ACCENT_BG = RGBColor(0xED, 0xF2, 0xF7)  # Gris clair (#EDF2F7)
C_RED_LIGHT = RGBColor(0xFF, 0xF5, 0xF5)  # Fond rouge léger (#FFF5F5)
C_GOLD = RGBColor(0xD9, 0x77, 0x06)       # Ambre doré (#D97706)

RUNNING_TITLE = "Conception d'un assistant conversationnel RAG en bancassurance vie : cas de SUNU Bank Togo"

def set_slide_background(slide, color=C_LIGHT_BG):
    """Définit la couleur d'arrière-plan d'une slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category_text=""):
    """Ajoute l'en-tête standard de la diapositive (Style Emmanuelle)."""
    # Titre principal en haut à gauche
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    
    if category_text:
        p_sub = tf.add_paragraph()
        p_sub.text = category_text
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = C_RED

def add_footer(slide, slide_num):
    """Ajoute le pied de page institutionnel avec le titre courant et le numéro de slide."""
    # Ligne fine de séparation
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.73), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = C_BORDER
    line.line.color.rgb = C_BORDER
    
    # Texte gauche
    tb_left = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(10.5), Inches(0.4))
    tf_left = tb_left.text_frame
    tf_left.margin_left = tf_left.margin_top = tf_left.margin_right = tf_left.margin_bottom = 0
    p_l = tf_left.paragraphs[0]
    p_l.text = RUNNING_TITLE
    p_l.font.name = "Segoe UI"
    p_l.font.size = Pt(9)
    p_l.font.italic = True
    p_l.font.color.rgb = C_MUTED
    
    # Numéro de slide à droite
    tb_right = slide.shapes.add_textbox(Inches(11.8), Inches(6.95), Inches(0.7), Inches(0.4))
    tf_right = tb_right.text_frame
    tf_right.margin_left = tf_right.margin_top = tf_right.margin_right = tf_right.margin_bottom = 0
    p_r = tf_right.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = str(slide_num)
    p_r.font.name = "Segoe UI"
    p_r.font.size = Pt(10)
    p_r.font.bold = True
    p_r.font.color.rgb = C_NAVY

def add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_BORDER, radius=None):
    """Ajoute une carte conteneur avec bordure propre."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def add_section_divider_slide(prs, section_num, section_title, section_subtitle=""):
    """Crée une slide intercalaire de section (Fond Bleu Nuit Foncé, Style Emmanuelle)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, C_NAVY)
    
    # Badge numéro de section
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(2.2), Inches(1.2), Inches(1.2))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_RED
    badge.line.color.rgb = C_WHITE
    badge.line.width = Pt(2)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    p_b.text = str(section_num)
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(28)
    p_b.font.bold = True
    p_b.font.color.rgb = C_WHITE
    
    # Titre de section
    tb = slide.shapes.add_textbox(Inches(2.8), Inches(2.1), Inches(9.2), Inches(1.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = section_title
    p.font.name = "Segoe UI"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    
    if section_subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = section_subtitle
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = RGBColor(0x90, 0xCD, 0xF4) # Bleu clair pastel
        
    # Ligne décorative en bas
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(4.5), Inches(10.8), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = C_RED
    line.line.color.rgb = C_RED
    
    # Mention bas de page
    tb_bot = slide.shapes.add_textbox(Inches(1.2), Inches(4.8), Inches(10.8), Inches(0.6))
    tf_bot = tb_bot.text_frame
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "Soutenance de Mémoire de Master • Collège de Paris / Université de Lomé • SUNU Bank Togo"
    p_bot.font.name = "Segoe UI"
    p_bot.font.size = Pt(11)
    p_bot.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    return slide

def build_presentation():
    """Génère la présentation PowerPoint complète (25 diapositives)."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    print("=== DÉBUT DE LA CRÉATION DU POWERPOINT DE SOUTENANCE ===")
    
    # ----------------------------------------------------
    # SLIDE 1 : PAGE DE GARDE (Title Slide)
    # ----------------------------------------------------
    print("Slide 1 : Page de Garde...")
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s1, C_WHITE)
    
    # Logos en haut
    logo_left = "Memoire/adogli_assets/cover_logo_1.png"
    if os.path.exists(logo_left):
        s1.shapes.add_picture(logo_left, Inches(0.8), Inches(0.5), width=Inches(2.5))
        
    logo_right = "Projet/LOGO-SUNU.png"
    if os.path.exists(logo_right):
        s1.shapes.add_picture(logo_right, Inches(10.2), Inches(0.4), width=Inches(2.2))
        
    # Diplôme
    tb_dip = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.73), Inches(0.6))
    tf_dip = tb_dip.text_frame
    p_dip = tf_dip.paragraphs[0]
    p_dip.alignment = PP_ALIGN.CENTER
    p_dip.text = "Mémoire en vue de l'obtention du diplôme de : Master en Intelligence Artificielle & Big Data"
    p_dip.font.name = "Segoe UI"
    p_dip.font.size = Pt(13)
    p_dip.font.bold = True
    p_dip.font.color.rgb = C_NAVY
    
    # Cadre du Titre du Mémoire
    card_title = add_card(s1, Inches(0.8), Inches(2.5), Inches(11.73), Inches(2.0), bg_color=C_LIGHT_BG, border_color=C_NAVY)
    tf_t = card_title.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.alignment = PP_ALIGN.CENTER
    p_t.text = "CONCEPTION ET IMPLÉMENTATION D'UN SYSTÈME RAG POUR L'AUTOMATISATION DE L'INFORMATION PRÉCONTRACTUELLE EN BANCASSURANCE VIE"
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(18)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY
    
    p_t_sub = tf_t.add_paragraph()
    p_t_sub.alignment = PP_ALIGN.CENTER
    p_t_sub.text = "Cas d'application : Portefeuille d'Assurance Vie de SUNU Bank Togo"
    p_t_sub.font.name = "Segoe UI"
    p_t_sub.font.size = Pt(14)
    p_t_sub.font.bold = True
    p_t_sub.font.color.rgb = C_RED
    
    # Présenté par & Encadreurs
    # Carte Candidat
    card_cand = add_card(s1, Inches(0.8), Inches(4.7), Inches(3.6), Inches(1.8), bg_color=C_WHITE, border_color=C_BORDER)
    tf_c = card_cand.text_frame
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "Présenté par :"
    p_c1.font.name = "Segoe UI"
    p_c1.font.size = Pt(11)
    p_c1.font.bold = True
    p_c1.font.color.rgb = C_RED
    p_c2 = tf_c.add_paragraph()
    p_c2.text = "ADOGLI Jean-Paul"
    p_c2.font.name = "Segoe UI"
    p_c2.font.size = Pt(13)
    p_c2.font.bold = True
    p_c2.font.color.rgb = C_DARK
    p_c3 = tf_c.add_paragraph()
    p_c3.text = "Option : Intelligence Artificielle & Big Data"
    p_c3.font.name = "Segoe UI"
    p_c3.font.size = Pt(9.5)
    p_c3.font.color.rgb = C_MUTED
    
    # Carte Encadreur Académique
    card_acad = add_card(s1, Inches(4.8), Inches(4.7), Inches(3.8), Inches(1.8), bg_color=C_WHITE, border_color=C_BORDER)
    tf_a = card_acad.text_frame
    p_a1 = tf_a.paragraphs[0]
    p_a1.text = "Encadreur Académique :"
    p_a1.font.name = "Segoe UI"
    p_a1.font.size = Pt(11)
    p_a1.font.bold = True
    p_a1.font.color.rgb = C_NAVY
    p_a2 = tf_a.add_paragraph()
    p_a2.text = "M. Latevi Sena LAWSON"
    p_a2.font.name = "Segoe UI"
    p_a2.font.size = Pt(12)
    p_a2.font.bold = True
    p_a2.font.color.rgb = C_DARK
    p_a3 = tf_a.add_paragraph()
    p_a3.text = "Ing. spécialiste webmaster\nCollège de Paris Supérieur / Université de Lomé"
    p_a3.font.name = "Segoe UI"
    p_a3.font.size = Pt(9.5)
    p_a3.font.color.rgb = C_MUTED
    
    # Carte Encadreur Professionnel
    card_prof = add_card(s1, Inches(8.9), Inches(4.7), Inches(3.6), Inches(1.8), bg_color=C_WHITE, border_color=C_BORDER)
    tf_p = card_prof.text_frame
    p_p1 = tf_p.paragraphs[0]
    p_p1.text = "Encadreur de Stage :"
    p_p1.font.name = "Segoe UI"
    p_p1.font.size = Pt(11)
    p_p1.font.bold = True
    p_p1.font.color.rgb = C_NAVY
    p_p2 = tf_p.add_paragraph()
    p_p2.text = "M. Kokou AGBOKOU"
    p_p2.font.name = "Segoe UI"
    p_p2.font.size = Pt(12)
    p_p2.font.bold = True
    p_p2.font.color.rgb = C_DARK
    p_p3 = tf_p.add_paragraph()
    p_p3.text = "Technicien informatique\nSUNU Bank Togo"
    p_p3.font.name = "Segoe UI"
    p_p3.font.size = Pt(9.5)
    p_p3.font.color.rgb = C_MUTED
    
    # Bas de page
    tb_foot = s1.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.73), Inches(0.4))
    tf_foot = tb_foot.text_frame
    p_f = tf_foot.paragraphs[0]
    p_f.alignment = PP_ALIGN.CENTER
    p_f.text = "LOME, TOGO • Année Académique 2024 - 2025"
    p_f.font.name = "Segoe UI"
    p_f.font.size = Pt(10)
    p_f.font.bold = True
    p_f.font.color.rgb = C_MUTED

    # ----------------------------------------------------
    # SLIDE 2 : SOMMAIRE / AGENDA
    # ----------------------------------------------------
    print("Slide 2 : Sommaire...")
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s2, C_LIGHT_BG)
    add_header(s2, "SOMMAIRE GENERAL DE LA SOUTENANCE", "Plan d'intervention")
    add_footer(s2, 2)
    
    agenda_items = [
        ("1", "INTRODUCTION GENERALE", "Contexte, Problématique, Hypothèses et Objectifs de recherche"),
        ("2", "CADRE THEORIQUE ET CONCEPTUEL", "Bancassurance, Paradigme RAG, LLM et Gouvernance CIMA/IPDCP"),
        ("3", "METHODOLOGIE DE L'ETUDE", "Démarche mixte, Corpus documentaire et Métriques d'évaluation"),
        ("4", "CONCEPTION ET PIPELINE RAG", "Architecture 4 couches, Ingestion, Indexation ChromaDB et Prompts"),
        ("5", "RESULTATS EXPERIMENTAUX & DEMONSTRATION", "Calibration du chunking, Embeddings, RAGAS et Captures de l'App"),
        ("6", "PROPOSITIONS D'INTERVENTION", "Déploiement Core Banking, Faisabilité et Recommandations managériales"),
        ("7", "CONCLUSION & PERSPECTIVES", "Bilan de la recherche, Apports académiques et Axes futurs")
    ]
    
    # Disposition en 2 colonnes de cartes
    for idx, (num, title, desc) in enumerate(agenda_items):
        col = 0 if idx < 4 else 1
        row = idx if idx < 4 else idx - 4
        left = Inches(0.8) if col == 0 else Inches(6.9)
        top = Inches(1.5 + row * 1.3)
        width = Inches(5.6)
        height = Inches(1.15)
        
        card = add_card(s2, left, top, width, height, bg_color=C_WHITE)
        
        # Badge Numéro
        badge = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.15), top + Inches(0.15), Inches(0.8), Inches(0.8))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_NAVY if num != "5" else C_RED
        badge.line.color.rgb = C_NAVY
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = num
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(18)
        p_b.font.bold = True
        p_b.font.color.rgb = C_WHITE
        
        # Texte Titre & Desc
        tb = s2.shapes.add_textbox(left + Inches(1.1), top + Inches(0.12), width - Inches(1.2), height - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9)
        p2.font.color.rgb = C_MUTED

    # ----------------------------------------------------
    # SECTION 1 : INTRODUCTION GÉNÉRALE
    # ----------------------------------------------------
    print("Slide 3-7 : Section 1 Introduction...")
    add_section_divider_slide(prs, 1, "Introduction Générale", "Contexte général – Problématique – Hypothèses – Objectifs")
    
    # Slide 4 : Contexte Général & Problématique
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s4, C_LIGHT_BG)
    add_header(s4, "CONTEXTE GENERAL & PROBLEMATIQUE CENTRALE", "Cadre de l'étude SUNU Bank Togo")
    add_footer(s4, 4)
    
    c_ctx = add_card(s4, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.1), bg_color=C_WHITE)
    tf_ctx = c_ctx.text_frame
    tf_ctx.word_wrap = True
    p = tf_ctx.paragraphs[0]
    p.text = "Le Contexte de la Bancassurance au Togo"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    
    ctx_bullets = [
        "Canal stratégique de distribution : Synergie entre SUNU Bank et SUNU Assurances Vie pour l'inclusion financière.",
        "Portefeuille ciblé : Offres Visa Études, Visa Études Plus (prévoyance éducation) et Horizon Retraite (épargne longue).",
        "Asymétrie d'information : Clauses contractuelles denses, termes techniques (valeur de rachat, frais CIMA, délais de carence).",
        "Surcharge des conseillers : Forte sollicitation en agence pour des questions d'information répétitives."
    ]
    for b in ctx_bullets:
        p = tf_ctx.add_paragraph()
        p.text = "• " + b
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_DARK
        
    c_pb = add_card(s4, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=C_RED_LIGHT, border_color=C_RED)
    tf_pb = c_pb.text_frame
    tf_pb.word_wrap = True
    p = tf_pb.paragraphs[0]
    p.text = "La Problématique Centrale de Recherche"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_RED
    
    p_q = tf_pb.add_paragraph()
    p_q.text = "« Dans quelle mesure l'automatisation de l'information précontractuelle par un système RAG (Retrieval-Augmented Generation) permet-elle d'améliorer la compréhension des clients tout en garantissant la conformité réglementaire CIMA et la fiabilité documentaire chez SUNU Bank Togo ? »"
    p_q.font.name = "Segoe UI"
    p_q.font.size = Pt(12)
    p_q.font.bold = True
    p_q.font.italic = True
    p_q.font.color.rgb = C_NAVY
    
    p_sub = tf_pb.add_paragraph()
    p_sub.text = "\nEnjeux majeurs identifiés :"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(11)
    p_sub.font.bold = True
    
    pb_bullets = [
        "Éliminer les hallucinations des modèles génératifs standards.",
        "Respecter strictement l'obligation légale d'information (Art. 6 & 74 CIMA).",
        "Garantir la confidentialité des données bancaires (Loi IPDCP Togo)."
    ]
    for b in pb_bullets:
        p = tf_pb.add_paragraph()
        p.text = "→ " + b
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # Slide 5 : Questions Spécifiques
    s5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s5, C_LIGHT_BG)
    add_header(s5, "QUESTIONS SPECIFIQUES DE RECHERCHE", "Déclinaison de la problématique")
    add_footer(s5, 5)
    
    q_data = [
        ("1", "Diagnostic & Besoins d'Information", "Quels sont les freins majeurs à la compréhension précontractuelle des produits d'assurance vie chez les prospects de SUNU Bank Togo, et comment caractériser leurs besoins d'information récurrents ?"),
        ("2", "Conception & Calibration de l'Architecture RAG", "Comment concevoir et calibrer un pipeline RAG (stratégie de chunking, modèle d'embeddings, seuils de similarité) capable de restituer avec exactitude les clauses des contrats d'assurance vie ?"),
        ("3", "Conformité Réglementaire & Traçabilité", "Quels mécanismes de sécurité (citation dynamique des sources, détection hors périmètre, escalade conseiller) permettent d'assurer une conformité stricte au Code CIMA et aux exigences de l'IPDCP Togo ?"),
        ("4", "Évaluation Empirique & Acceptabilité", "Quels sont les impacts mesurables du prototype conversationnel sur l'exactitude des réponses (RAGAS), la réduction du temps de conseil et l'acceptation par les utilisateurs (TAM) ?")
    ]
    for idx, (num, title, text) in enumerate(q_data):
        top = Inches(1.5 + idx * 1.3)
        card = add_card(s5, Inches(0.8), top, Inches(11.73), Inches(1.15), bg_color=C_WHITE)
        
        # Badge
        badge = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), top + Inches(0.2), Inches(0.75), Inches(0.75))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_NAVY
        badge.line.color.rgb = C_NAVY
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = f"Q{num}"
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = C_WHITE
        
        tb = s5.shapes.add_textbox(Inches(1.9), top + Inches(0.1), Inches(10.4), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_RED
        
        p2 = tf.add_paragraph()
        p2.text = text
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_DARK

    # Slide 6 : Hypothèses de Recherche
    s6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s6, C_LIGHT_BG)
    add_header(s6, "HYPOTHESES DE L'ETUDE", "Cadre de validation empirique")
    add_footer(s6, 6)
    
    # Hypothèse générale
    c_hg = add_card(s6, Inches(0.8), Inches(1.4), Inches(11.73), Inches(1.2), bg_color=C_NAVY, border_color=C_NAVY)
    tf_hg = c_hg.text_frame
    tf_hg.word_wrap = True
    p = tf_hg.paragraphs[0]
    p.text = "HYPOTHÈSE GÉNÉRALE (HG)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x90, 0xCD, 0xF4)
    p_t = tf_hg.add_paragraph()
    p_t.text = "L'implémentation d'un assistant conversationnel RAG fondé sur un corpus documentaire certifié permet d'automatiser l'information précontractuelle en bancassurance vie, en offrant une exactitude documentaire supérieure à 75% et en réduisant les asymétries d'information."
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(10.5)
    p_t.font.color.rgb = C_WHITE
    
    h_cards = [
        ("H1", "Calibration du Retrieval", "Un découpage documentaire calibré (taille 300 / overlap 30) combiné à all-MiniLM-L6-v2 maximise le taux de récupération (Source Hit@5 ≥ 75%)."),
        ("H2", "Fidélité & Anti-Hallucination", "L'ancrage contextuel strict et l'injection de prompts de refus éliminent les hallucinations factuelles (Score Faithfulness RAGAS ≥ 0.80)."),
        ("H3", "Conformité CIMA & Confiance", "La citation explicite des clauses contractuelles et l'escalade conseiller renforcent la conformité légale et la confiance perçue (TAM ≥ 4/5).")
    ]
    for idx, (code, title, text) in enumerate(h_cards):
        left = Inches(0.8 + idx * 4.0)
        c = add_card(s6, left, Inches(2.8), Inches(3.75), Inches(3.8), bg_color=C_WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = code
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_RED
        
        p_tit = tf.add_paragraph()
        p_tit.text = title
        p_tit.font.name = "Segoe UI"
        p_tit.font.size = Pt(11.5)
        p_tit.font.bold = True
        p_tit.font.color.rgb = C_NAVY
        
        p_txt = tf.add_paragraph()
        p_txt.text = "\n" + text
        p_txt.font.name = "Segoe UI"
        p_txt.font.size = Pt(10)
        p_txt.font.color.rgb = C_DARK

    # Slide 7 : Objectifs de Recherche
    s7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s7, C_LIGHT_BG)
    add_header(s7, "OBJECTIFS DE L'ETUDE", "Finalités académiques et opérationnelles")
    add_footer(s7, 7)
    
    c_og = add_card(s7, Inches(0.8), Inches(1.4), Inches(11.73), Inches(1.1), bg_color=C_RED, border_color=C_RED)
    tf_og = c_og.text_frame
    tf_og.word_wrap = True
    p = tf_og.paragraphs[0]
    p.text = "OBJECTIF GÉNÉRAL"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p_t = tf_og.add_paragraph()
    p_t.text = "Concevoir, implémenter et évaluer empiriquement un système RAG certifié pour automatiser l'information précontractuelle en bancassurance vie chez SUNU Bank Togo."
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE
    
    obj_cards = [
        ("1", "Structuration du Corpus", "Collecter, auditer et segmenter les 5 documents sources officiels (spécifications, conditions générales, notice CIMA) en 150 chunks indexés."),
        ("2", "Développement du Pipeline", "Implémenter l'architecture LangChain/ChromaDB avec garde-fous CIMA (Articles 6 et 74), citation dynamique et module d'escalade."),
        ("3", "Évaluation Expérimentale", "Mesurer rigoureusement les métriques de retrieval (Hit@k, MRR), de génération (RAGAS) et tester le prototype sur 75 requêtes métiers."),
        ("4", "Stratégie de Déploiement", "Formuler un plan d'intégration opérationnel au Core Banking et aux canaux digitaux (Web, WhatsApp) de SUNU Bank Togo.")
    ]
    for idx, (num, title, text) in enumerate(obj_cards):
        row = 0 if idx < 2 else 1
        col = idx % 2
        left = Inches(0.8) if col == 0 else Inches(6.8)
        top = Inches(2.7 + row * 2.0)
        c = add_card(s7, left, top, Inches(5.7), Inches(1.85), bg_color=C_WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"Objectif Spécifique {num} : {title}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        
        p_txt = tf.add_paragraph()
        p_txt.text = text
        p_txt.font.name = "Segoe UI"
        p_txt.font.size = Pt(9.5)
        p_txt.font.color.rgb = C_DARK

    # ----------------------------------------------------
    # SECTION 2 : CADRE THÉORIQUE ET CONCEPTUEL
    # ----------------------------------------------------
    print("Slide 8-10 : Section 2 Cadre Théorique...")
    add_section_divider_slide(prs, 2, "Cadre Théorique et Conceptuel", "Concepts clés – Modèles théoriques – Cadre réglementaire CIMA / IPDCP")
    
    # Slide 9 : Concepts Clés
    s9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s9, C_LIGHT_BG)
    add_header(s9, "CONCEPTS CLES DU SUJET", "Fondements conceptuels et technologiques")
    add_footer(s9, 9)
    
    concepts = [
        ("BANCASSURANCE VIE", "Distribution de contrats d'assurance vie (prévoyance et épargne) à travers le réseau et la clientèle d'une institution bancaire (SUNU Bank Togo)."),
        ("RAG (Retrieval-Augmented Generation)", "Architecture d'IA combinant un moteur de recherche vectoriel (retriever) et un modèle de langage (LLM) pour générer des réponses ancrées sur des sources vérifiées."),
        ("INFORMATION PRECONTRACTUELLE", "Ensemble des renseignements obligatoires transmis au souscripteur avant l'engagement (garanties, cotisations, exclusions, valeur de rachat)."),
        ("GOUVERNANCE CIMA & IPDCP", "Cadre juridique supranational (Code CIMA Art. 6 & 74) et national (Loi IPDCP Togo n° 2019-014) régissant la loyauté commerciale et les données personnelles.")
    ]
    for idx, (title, desc) in enumerate(concepts):
        top = Inches(1.5 + idx * 1.3)
        c = add_card(s9, Inches(0.8), top, Inches(11.73), Inches(1.15), bg_color=C_WHITE)
        
        # Label gauche
        tb_l = s9.shapes.add_textbox(Inches(1.0), top + Inches(0.2), Inches(3.4), Inches(0.75))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        p_l = tf_l.paragraphs[0]
        p_l.text = title
        p_l.font.name = "Segoe UI"
        p_l.font.size = Pt(11)
        p_l.font.bold = True
        p_l.font.color.rgb = C_NAVY if idx % 2 == 0 else C_RED
        
        # Ligne verticale
        line = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), top + Inches(0.15), Inches(0.02), Inches(0.85))
        line.fill.solid()
        line.fill.fore_color.rgb = C_BORDER
        line.line.color.rgb = C_BORDER
        
        # Description droite
        tb_r = s9.shapes.add_textbox(Inches(4.7), top + Inches(0.15), Inches(7.6), Inches(0.85))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        p_r = tf_r.paragraphs[0]
        p_r.text = desc
        p_r.font.name = "Segoe UI"
        p_r.font.size = Pt(10)
        p_r.font.color.rgb = C_DARK

    # Slide 10 : Modèles Théoriques
    s10 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s10, C_LIGHT_BG)
    add_header(s10, "MODELES THEORIQUES DE REFERENCE", "Ancrage scientifique de la recherche")
    add_footer(s10, 10)
    
    theories = [
        ("Modèle d'Acceptation de la Technologie (TAM)", "Davis (1989), Venkatesh & Davis (2000)", "Analyse de l'utilité perçue (PU), de la facilité d'utilisation perçue (PEOU) et de la confiance dans l'adoption d'un assistant conversationnel bancaire par les prospects."),
        ("Paradigme RAG & Bi-Encodeurs Sémantiques", "Lewis et al. (2020), Reimers & Gurevych (2019)", "Modélisation de la recherche d'information dense par projection vectorielle dans un espace latent (embeddings) et injection de contexte documentaire certifié."),
        ("Théorie de l'Asymétrie d'Information", "Akerlof (1970), Stiglitz & Weiss (1981)", "Explication du risque d'incompréhension et de sélection adverse en assurance vie, et justification de la vulgarisation interactive pour rétablir la transparence.")
    ]
    for idx, (nom, auteurs, detail) in enumerate(theories):
        left = Inches(0.8 + idx * 4.0)
        c = add_card(s10, left, Inches(1.5), Inches(3.75), Inches(5.1), bg_color=C_WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = nom
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        
        p_aut = tf.add_paragraph()
        p_aut.text = auteurs
        p_aut.font.name = "Segoe UI"
        p_aut.font.size = Pt(9.5)
        p_aut.font.bold = True
        p_aut.font.color.rgb = C_RED
        
        p_det = tf.add_paragraph()
        p_det.text = "\n" + detail
        p_det.font.name = "Segoe UI"
        p_det.font.size = Pt(10)
        p_det.font.color.rgb = C_DARK

    # ----------------------------------------------------
    # SECTION 3 : MÉTHODOLOGIE DE L'ÉTUDE
    # ----------------------------------------------------
    print("Slide 11-13 : Section 3 Méthodologie...")
    add_section_divider_slide(prs, 3, "Méthodologie de l'Étude", "Démarche de recherche – Corpus documentaire – Protocole expérimental")
    
    # Slide 12 : Démarche & Corpus Documentaire
    s12 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s12, C_LIGHT_BG)
    add_header(s12, "DEMARCHE METHODOLOGIQUE & CORPUS D'ETUDE", "Design de recherche mixte")
    add_footer(s12, 12)
    
    # Gauche : Démarche
    c_m = add_card(s12, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.1), bg_color=C_WHITE)
    tf_m = c_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "Démarche Méthodologique Mixte"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    
    m_steps = [
        "1. Recherche appliquée & ingénierie : Développement d'un prototype conversationnel RAG connecté à ChromaDB.",
        "2. Corpus documentaire certifié : 5 documents officiels de SUNU Bank Togo (150 chunks indexés).",
        "3. Jeu d'évaluation étalonné : 75 questions réparties sur 11 thématiques précontractuelles et pièges hors périmètre.",
        "4. Approche empirique : Benchmark de 6 configurations de chunking et comparaison de 2 modèles d'embeddings."
    ]
    for s in m_steps:
        p = tf_m.add_paragraph()
        p.text = s
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_DARK
        
    # Droite : Corpus Image (Figure III.1)
    c_fig = add_card(s12, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=C_WHITE)
    tf_f = c_fig.text_frame
    p = tf_f.paragraphs[0]
    p.text = "Volumétrie du Corpus (150 Chunks)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    
    img_corpus = "Memoire/figures/fig_corpus_distribution.png"
    if os.path.exists(img_corpus):
        s12.shapes.add_picture(img_corpus, Inches(7.0), Inches(2.1), width=Inches(5.3))

    # Slide 13 : Variables & Métriques d'Évaluation
    s13 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s13, C_LIGHT_BG)
    add_header(s13, "VARIABLES ET METRIQUES D'EVALUATION", "Protocole quantitatif et qualitatif")
    add_footer(s13, 13)
    
    metrics = [
        ("MÉTRIQUES DE RETRIEVAL (IR)", [
            "Source Hit@5 : Pourcentage où le document officiel attendu figure dans le top-5 (Cible ≥ 75%).",
            "Mean Reciprocal Rank (MRR) : Rang moyen réciproque du premier chunk pertinent.",
            "Precision@1 : Exactitude du tout premier passage récupéré."
        ]),
        ("MÉTRIQUES DE GÉNÉRATION (RAGAS)", [
            "Faithfulness (Fidélité) : Vérification que la réponse découle uniquement du contexte (zéro hallucination).",
            "Answer Relevance (Pertinence) : Adéquation directe de la réponse à la question posée.",
            "Context Precision & Recall : Qualité du filtrage documentaire."
        ]),
        ("MÉTRIQUES D'ACCEPTATION (TAM)", [
            "Utilité perçue (PU) : Clarté et vulgarisation des garanties (Échelle 1-5).",
            "Facilité d'utilisation (PEOU) : Fluidité de l'interface conversationnelle.",
            "Confiance perçue : Traçabilité des citations d'articles CIMA."
        ])
    ]
    for idx, (title, items) in enumerate(metrics):
        left = Inches(0.8 + idx * 4.0)
        c = add_card(s13, left, Inches(1.5), Inches(3.75), Inches(5.1), bg_color=C_WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        
        for it in items:
            p_it = tf.add_paragraph()
            p_it.text = "• " + it
            p_it.font.name = "Segoe UI"
            p_it.font.size = Pt(9.5)
            p_it.font.color.rgb = C_DARK

    # ----------------------------------------------------
    # SECTION 4 : CONCEPTION ET PIPELINE RAG
    # ----------------------------------------------------
    print("Slide 14-16 : Section 4 Architecture & Pipeline...")
    add_section_divider_slide(prs, 4, "Conception et Pipeline RAG", "Architecture modulaire 4 couches – Flux séquentiel – Garde-fous CIMA")
    
    # Slide 15 : Architecture Modulaire (Figure I.1)
    s15 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s15, C_LIGHT_BG)
    add_header(s15, "ARCHITECTURE MODULAIRE EN 4 COUCHES", "Figure I.1 du Mémoire")
    add_footer(s15, 15)
    
    img_arch = "Memoire/figures/fig_architecture.png"
    if os.path.exists(img_arch):
        s15.shapes.add_picture(img_arch, Inches(0.8), Inches(1.5), width=Inches(11.73))

    # Slide 16 : Pipeline Séquentiel et Escalade (Figure I.2)
    s16 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s16, C_LIGHT_BG)
    add_header(s16, "PIPELINE SEQUENTIEL & ESCALADE CONSEILLER", "Figure I.2 du Mémoire")
    add_footer(s16, 16)
    
    img_pipe = "Memoire/figures/fig_pipeline.png"
    if os.path.exists(img_pipe):
        s16.shapes.add_picture(img_pipe, Inches(0.8), Inches(1.5), width=Inches(11.73))

    # ----------------------------------------------------
    # SECTION 5 : RÉSULTATS EXPÉRIMENTAUX & DÉMONSTRATION
    # ----------------------------------------------------
    print("Slide 17-19 : Section 5 Résultats & Démonstration...")
    add_section_divider_slide(prs, 5, "Résultats Expérimentaux & Démonstration", "Calibration du chunking – Embeddings – Captures réelles de l'application")
    
    # Slide 18 : Calibration Chunking & Embeddings
    s18 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s18, C_LIGHT_BG)
    add_header(s18, "RESULTATS DE CALIBRATION EMPIRIQUE", "Chunking (Fig IV.2) et Embeddings (Fig IV.3)")
    add_footer(s18, 18)
    
    # Image Chunking gauche
    img_chk = "Memoire/figures/fig_chunking.png"
    if os.path.exists(img_chk):
        s18.shapes.add_picture(img_chk, Inches(0.8), Inches(1.5), width=Inches(5.7))
        
    # Image Embeddings droite
    img_emb = "Memoire/figures/fig_embeddings.png"
    if os.path.exists(img_emb):
        s18.shapes.add_picture(img_emb, Inches(6.8), Inches(1.5), width=Inches(5.7))
        
    # Synthèse bas
    c_syn = add_card(s18, Inches(0.8), Inches(5.6), Inches(11.73), Inches(1.1), bg_color=C_WHITE)
    tf_s = c_syn.text_frame
    p = tf_s.paragraphs[0]
    p.text = "Enseignements majeurs de l'expérimentation :"
    p.font.name = "Segoe UI"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p_t = tf_s.add_paragraph()
    p_t.text = "• Configuration optimale : Taille de chunk de 300 caractères avec chevauchement de 30 (Hit@5 = 78,7 % | MRR = 0,434).\n• Modèle d'embeddings : all-MiniLM-L6-v2 surpasse le modèle multilingue en précision sur le vocabulaire technique financier."
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(9.5)
    p_t.font.color.rgb = C_DARK

    # Slide 19 : Démonstration - Captures Réelles de l'Application
    s19 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s19, C_LIGHT_BG)
    add_header(s19, "DEMONSTRATION DU PROTOTYPE APPLICATIF", "Captures originales de l'interface SUNU Bank")
    add_footer(s19, 19)
    
    # Capture 1 : Accueil Sombre (01_accueil_concierge_sombre.png)
    img_ui1 = "Projet/docs/screenshots/01_accueil_concierge_sombre.png"
    if os.path.exists(img_ui1):
        s19.shapes.add_picture(img_ui1, Inches(0.8), Inches(1.5), width=Inches(3.7))
        
    # Capture 2 : Discussion IA (03_chat_discussion_ia.png)
    img_ui2 = "Projet/docs/screenshots/03_chat_discussion_ia.png"
    if os.path.exists(img_ui2):
        s19.shapes.add_picture(img_ui2, Inches(4.8), Inches(1.5), width=Inches(3.7))
        
    # Capture 3 : Citations Sources (04_reponse_concierge_visa_etudes.png)
    img_ui3 = "Projet/docs/screenshots/04_reponse_concierge_visa_etudes.png"
    if os.path.exists(img_ui3):
        s19.shapes.add_picture(img_ui3, Inches(8.8), Inches(1.5), width=Inches(3.7))
        
    # Légendes sous les 3 images
    captions = [
        ("Accueil Concierge (Mode Sombre)", Inches(0.8)),
        ("Discussion Conversationnelle IA", Inches(4.8)),
        ("Réponse & Citation Code CIMA", Inches(8.8))
    ]
    for cap, left_pos in captions:
        tb = s19.shapes.add_textbox(left_pos, Inches(4.0), Inches(3.7), Inches(0.5))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = cap
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_NAVY

    # Cartouche de conformité en bas
    c_conf = add_card(s19, Inches(0.8), Inches(4.6), Inches(11.73), Inches(2.1), bg_color=C_WHITE)
    tf_c = c_conf.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Fonctionnalités Clés Démontrées sur le Prototype Réel :"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_RED
    
    feats = [
        "• Citation dynamique et transparente des documents sources (Conditions Générales, Fiches Produits, Notice CIMA Art. 74).",
        "• Rejet automatique et poli des requêtes hors domaine avec proposition d'escalade vers un conseiller agence.",
        "• Sécurité des données : Environnement chiffré SSL/TLS et conformité intégrale aux exigences de l'IPDCP Togo."
    ]
    for f in feats:
        p = tf_c.add_paragraph()
        p.text = f
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_DARK

    # ----------------------------------------------------
    # SECTION 6 : PROPOSITIONS D'INTERVENTION
    # ----------------------------------------------------
    print("Slide 20-22 : Section 6 Propositions & Faisabilité...")
    add_section_divider_slide(prs, 6, "Propositions d'Intervention & Déploiement", "Plan d'intégration opérationnelle – Faisabilité – Recommandations")
    
    # Slide 21 : Propositions d'Intervention
    s21 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s21, C_LIGHT_BG)
    add_header(s21, "RECOMMANDATIONS & PLAN D'INTERVENTION", "Stratégie de déploiement chez SUNU Bank Togo")
    add_footer(s21, 21)
    
    props = [
        ("1", "Intégration Core Banking & CRM", "Connecter l'assistant via API sécurisée au portail bancaire pour pré-remplir les formulaires de souscription et historiser les demandes d'escalade."),
        ("2", "Déploiement Omnicanal", "Étendre la solution au-delà du Web vers WhatsApp Business et les bornes interactives dans les 15 agences SUNU Bank Togo."),
        ("3", "Comité de Gouvernance du Corpus", "Mettre en place une revue trimestrielle du corpus vectoriel associant juristes, actuaires et conseillers pour actualiser les clauses."),
        ("4", "Formation & Conduite du Changement", "Former les conseillers clientèle à exploiter l'assistant comme outil d'aide à la vente pour consacrer plus de temps au conseil à forte valeur ajoutée.")
    ]
    for idx, (num, title, desc) in enumerate(props):
        top = Inches(1.5 + idx * 1.3)
        c = add_card(s21, Inches(0.8), top, Inches(11.73), Inches(1.15), bg_color=C_WHITE)
        
        # Badge
        badge = s21.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top + Inches(0.18), Inches(0.8), Inches(0.8))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_NAVY
        badge.line.color.rgb = C_NAVY
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = num
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(16)
        p_b.font.bold = True
        p_b.font.color.rgb = C_WHITE
        
        tb = s21.shapes.add_textbox(Inches(2.0), top + Inches(0.12), Inches(10.3), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = C_RED
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_DARK

    # Slide 22 : Faisabilité
    s22 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s22, C_LIGHT_BG)
    add_header(s22, "ANALYSE DE FAISABILITE MULTIDIMENSIONNELLE", "Garanties de viabilité du projet")
    add_footer(s22, 22)
    
    faisabilites = [
        ("FAISABILITÉ TECHNIQUE", "Stack logicielle éprouvée et modulaire (Python, LangChain, ChromaDB, FastAPI/React). Possibilité d'hébergement hybride (Cloud sécurisé ou On-Premise) avec latence d'inférence < 2 secondes."),
        ("FAISABILITÉ ÉCONOMIQUE", "Investissement modéré basé sur des technologies open-source. ROI rapide estimé grâce au désengorgement des agences et à l'accélération du taux de conversion précontractuel."),
        ("FAISABILITÉ JURIDIQUE", "Conformité native avec le Code CIMA (Articles 6 et 74) grâce aux citations contractuelles obligatoires et au respect de la Loi n° 2019-014 (IPDCP Togo) sur la vie privée.")
    ]
    for idx, (title, text) in enumerate(faisabilites):
        left = Inches(0.8 + idx * 4.0)
        c = add_card(s22, left, Inches(1.5), Inches(3.75), Inches(5.1), bg_color=C_WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        
        p_txt = tf.add_paragraph()
        p_txt.text = "\n" + text
        p_txt.font.name = "Segoe UI"
        p_txt.font.size = Pt(10)
        p_txt.font.color.rgb = C_DARK

    # Slide 23 : Apports de l'Étude (Plan Scientifique, Managérial, Méthodologique)
    s23 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s23, C_LIGHT_BG)
    add_header(s23, "APPORTS MAJEURS DE L'ETUDE", "Contributions scientifiques et professionnelles")
    add_footer(s23, 23)
    
    apports = [
        ("PLAN SCIENTIFIQUE", [
            "Validation empirique de l'architecture RAG sur un corpus réglementaire financier en langue française.",
            "Modélisation de la réduction des asymétries d'information par interaction conversationnelle ancrée."
        ]),
        ("PLAN MANAGÉRIAL & BANCAIRE", [
            "Mise à disposition d'un prototype fonctionnel prêt au déploiement pour SUNU Bank Togo.",
            "Sécurisation du parcours client et réduction du risque de litige précontractuel (Code CIMA)."
        ]),
        ("PLAN MÉTHODOLOGIQUE", [
            "Protocole reproductible de calibration du chunking et des embeddings pour les institutions financières.",
            "Grille d'évaluation combinant métriques de Recherche d'Information (IR), RAGAS et TAM."
        ])
    ]
    for idx, (title, items) in enumerate(apports):
        top = Inches(1.5 + idx * 1.7)
        c = add_card(s23, Inches(0.8), top, Inches(11.73), Inches(1.55), bg_color=C_WHITE)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_RED if idx == 1 else C_NAVY
        
        for it in items:
            p_it = tf.add_paragraph()
            p_it.text = "• " + it
            p_it.font.name = "Segoe UI"
            p_it.font.size = Pt(9.5)
            p_it.font.color.rgb = C_DARK

    # ----------------------------------------------------
    # SECTION 7 : CONCLUSION & DIAPOSITIVE DE FIN
    # ----------------------------------------------------
    print("Slide 24-25 : Conclusion & Fin...")
    
    # Slide 24 : Conclusion Générale & Perspectives
    s24 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s24, C_LIGHT_BG)
    add_header(s24, "CONCLUSION GENERALE & PERSPECTIVES", "Bilan des travaux de recherche")
    add_footer(s24, 24)
    
    c_cgl = add_card(s24, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.1), bg_color=C_WHITE)
    tf_cgl = c_cgl.text_frame
    tf_cgl.word_wrap = True
    p = tf_cgl.paragraphs[0]
    p.text = "Synthèse des Résultats"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    
    cgl_bullets = [
        "Démonstration concluante : Le système RAG résout efficacement le dilemme entre automatisation et conformité documentaire.",
        "Performance validée : 78,7 % de Hit@5 et fidélité factuelle RAGAS de 0,84 sur les 75 requêtes tests.",
        "Impact opérationnel : Réduction du temps de renseignement préalable et meilleure traçabilité des conseils fournis."
    ]
    for b in cgl_bullets:
        p = tf_cgl.add_paragraph()
        p.text = "• " + b
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK
        
    c_persp = add_card(s24, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=C_WHITE)
    tf_persp = c_persp.text_frame
    tf_persp.word_wrap = True
    p = tf_persp.paragraphs[0]
    p.text = "Perspectives & Travaux Futurs"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_RED
    
    persp_bullets = [
        "Extension du périmètre : Intégrer les produits bancaires purs (crédits immobiliers, épargne, cartes monétiques).",
        "Modèles vocaux & multilinguisme : Développer des interfaces vocales en langues locales (Éwé, Kabyè) pour renforcer l'inclusion financière.",
        "Fine-tuning localisé : Entraîner un petit modèle spécialisé (SLM) sur le corpus juridique UEMOA/CIMA."
    ]
    for b in persp_bullets:
        p = tf_persp.add_paragraph()
        p.text = "→ " + b
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # Slide 25 : Diapositive de Fin (Merci)
    s25 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(s25, C_NAVY)
    
    # Carte centrale élégante
    c_end = add_card(s25, Inches(2.0), Inches(1.5), Inches(9.33), Inches(4.5), bg_color=C_NAVY, border_color=C_RED)
    tf_end = c_end.text_frame
    tf_end.word_wrap = True
    
    p = tf_end.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "MERCI POUR VOTRE ATTENTION !"
    p.font.name = "Segoe UI"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    
    p_sub = tf_end.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = "\nLa parole est désormais aux membres du Jury pour les questions et échanges."
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(0x90, 0xCD, 0xF4)
    
    p_author = tf_end.add_paragraph()
    p_author.alignment = PP_ALIGN.CENTER
    p_author.text = "\nADOGLI Jean-Paul • Master en Intelligence Artificielle & Big Data\nCollège de Paris Supérieur / Université de Lomé • SUNU Bank Togo"
    p_author.font.name = "Segoe UI"
    p_author.font.size = Pt(11)
    p_author.font.color.rgb = C_MUTED

    prs.save(PPTX_OUTPUT)
    print(f"=== [SUCCÈS] Présentation générée avec succès : {PPTX_OUTPUT} ({len(prs.slides)} diapositives) ===")

if __name__ == "__main__":
    build_presentation()
